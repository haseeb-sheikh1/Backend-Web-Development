"""
generate_test_files.py
Generates test_students.docx and test_students.pdf with messy multi-student notes.
Run once: python generate_test_files.py
"""

# ── DOCX ─────────────────────────────────────────────────────────────────────
from docx import Document

doc = Document()
doc.add_heading("Admissions Office — Rough Interview Notes", 0)

doc.add_paragraph(
    "Meeting notes from 22nd Oct — staff: Mr. Tariq\n"
    "These are informal, please clean up before filing.\n"
)

doc.add_heading("Student 1", level=2)
doc.add_paragraph(
    "Sara Khan came in at 9am. She's very keen on BS Artificial Intelligence. "
    "Her FSc marks were 89 percent. Contact: sara.khan@outlook.com, "
    "phone number is 0311-2233445."
)

doc.add_heading("Student 2", level=2)
doc.add_paragraph(
    "Next was Kamran Mirza — no appointment, just walked in. "
    "Interested in MBA program. Graduated with 74%. "
    "He left his email: kamran.mirza92@gmail.com. Didn't share phone."
)

doc.add_heading("Student 3", level=2)
doc.add_paragraph(
    "Nadia Hussain — BS Data Science. Scored 91.5% in intermediate. "
    "nadia.h@yahoo.com | ph: 0300-8877665"
)

doc.add_heading("Student 4", level=2)
doc.add_paragraph(
    "Omar Farooq. Program of interest: MBBS. "
    "marks = 82 percent. email not provided. cell: 0345-6543210"
)

doc.add_heading("Student 5", level=2)
doc.add_paragraph(
    "Hina Baig visited after lunch. Wants to study BS Psychology. "
    "She mentioned scoring around 77% in her exams. "
    "hina.baig2024@hotmail.com, no phone given."
)

doc.add_paragraph("\n-- End of notes for today --")

doc.save("test_students.docx")
print("  [OK] test_students.docx created")


# ── PDF ──────────────────────────────────────────────────────────────────────
from fpdf import FPDF

pdf = FPDF()
pdf.set_auto_page_break(auto=True, margin=15)
pdf.add_page()

pdf.set_font("Helvetica", "B", 16)
pdf.cell(0, 10, "University Admission Notes - Oct 2026", ln=True)
pdf.set_font("Helvetica", "", 10)
pdf.cell(0, 6, "Compiled by: Admissions Coordinator", ln=True)
pdf.ln(4)

students = [
    {
        "title": "Applicant: Fatima Malik",
        "body": (
            "Fatima came in with her father. She wants BS Computer Science. "
            "Her intermediate result shows 93 percent. "
            "Email: fatima.malik@fast.edu.pk, Phone: 0321-4455667"
        ),
    },
    {
        "title": "Walk-in: Umar Shahid",
        "body": (
            "Umar Shahid — interested in BS Electrical Engineering. "
            "He scored 88% overall. Provided contact: umar.shahid@live.com "
            "Mobile: 0333-7788990"
        ),
    },
    {
        "title": "Email inquiry — Zainab Iqbal",
        "body": (
            "Received email from zainab.iqbal@gmail.com. "
            "Wants to enroll in BBA. Her marks: 79 percent. "
            "No phone number shared in the email."
        ),
    },
    {
        "title": "Referred by faculty: Asad Mehmood",
        "body": (
            "Asad Mehmood referred by Dr. Khalid. "
            "Program: BS Mathematics. Scored 85.5% in FSc. "
            "asad.mehmood@hotmail.com | 0312-9900112"
        ),
    },
    {
        "title": "Late arrival: Mariam Siddiqui",
        "body": (
            "Mariam arrived just before closing. Interested in MS Data Science. "
            "Undergrad marks: 72 percent. "
            "mariam.siddiqui@outlook.com, 0300-1122334"
        ),
    },
]

for s in students:
    pdf.set_font("Helvetica", "B", 12)
    pdf.cell(0, 8, s["title"], ln=True)
    pdf.set_font("Helvetica", "", 10)
    pdf.multi_cell(0, 6, s["body"])
    pdf.ln(3)
    pdf.set_draw_color(200, 200, 200)
    pdf.line(10, pdf.get_y(), 200, pdf.get_y())
    pdf.ln(4)

pdf.output("test_students.pdf")
print("  [OK] test_students.pdf created")

print("\n  Both test files are ready in your AGENT folder!")
